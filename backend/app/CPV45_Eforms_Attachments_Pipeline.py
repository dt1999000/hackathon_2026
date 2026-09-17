# ruff: noqa: T201
"""CPV 45 eForms and Tender Documents Pipeline.

Downloads the CSV and eForms exports, filters construction lots whose CPV
starts with 45, parses notice XML, and writes agent_input.jsonl.

Run:
    python CPV45_Eforms_Attachments_Pipeline.py

Paths are resolved from this file's directory, so the working directory
does not matter. Configure behaviour with CPV45_* environment variables.
"""

from __future__ import annotations


# CPV 45 eForms and Tender Documents Pipeline
#
# This notebook builds the evidence package for a construction-procurement agent:
#
# 1. download the CSV and eForms exports for the same publication day or month;
# 2. filter the CSV tables to CPV codes beginning with `45`;
# 3. match the selected notice versions and lots to their original eForms XML;
# 4. extract deadlines, qualification criteria, guarantees, execution conditions and document URLs;
# 5. optionally download direct attachment files and queue platform landing pages for an adapter or manual review;
# 6. extract text from available files and create one JSON object per lot for the downstream agent.
#
# The OpenData ZIP contains notices, not the complete tender-document package. Document links normally lead to the original procurement platform.


# 1. Configuration
#
# The default source is yesterday's daily API export because the API does not publish the current day. Set `SOURCE_MODE` to `local` to use previously downloaded ZIP files.
#
# Attachment downloading is intentionally disabled by default. First inspect `document_links.csv`. Enable it only after confirming that the relevant platforms allow automated retrieval and after setting a descriptive user agent.


import hashlib
import io
import json
import mimetypes
import os
import re
import shutil
import ssl
import time
import zipfile
from collections import defaultdict
from datetime import date, timedelta
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, Iterable
from urllib.parse import unquote, urlencode, urljoin, urlparse
from urllib.request import Request, urlopen
import xml.etree.ElementTree as ET

import certifi
import pandas as pd

SCRIPT_DIRECTORY = Path(__file__).resolve().parent
os.chdir(SCRIPT_DIRECTORY)


def display(value):
    if isinstance(value, pd.DataFrame):
        print(value.to_string(index=False))
        return
    print(value)


SSL_CONTEXT = ssl.create_default_context(cafile=certifi.where())

API_URL = 'https://oeffentlichevergabe.de/api/notice-exports'


def env_bool(name: str, default: bool) -> bool:
    value = os.getenv(name)
    if value is None:
        return default
    return value.strip().lower() in {'1', 'true', 'yes', 'on'}


# api | local
SOURCE_MODE = os.getenv('CPV45_SOURCE_MODE', 'api').lower()

# day: YYYY-MM-DD, month: YYYY-MM
PERIOD_TYPE = os.getenv('CPV45_PERIOD_TYPE', 'day').lower()
PERIOD_VALUE = os.getenv(
    'CPV45_PERIOD_VALUE',
    (date.today() - timedelta(days=1)).isoformat(),
)

LOCAL_CSV_ZIP = Path(os.getenv('CPV45_LOCAL_CSV_ZIP', 'data/notices_csv.zip'))
LOCAL_EFORMS_ZIP = Path(os.getenv('CPV45_LOCAL_EFORMS_ZIP', 'data/notices_eforms.zip'))

CACHE_DIRECTORY = Path(os.getenv('CPV45_CACHE_DIRECTORY', 'data/cache'))
OUTPUT_DIRECTORY = Path(os.getenv('CPV45_OUTPUT_DIRECTORY', 'outputs/cpv45_pipeline'))
ATTACHMENT_DIRECTORY = OUTPUT_DIRECTORY / 'attachments'
MANUAL_ATTACHMENTS_DIRECTORY = Path(
    os.getenv('CPV45_MANUAL_ATTACHMENTS', 'data/manual_attachments')
)

CPV_PREFIX = os.getenv('CPV45_PREFIX', '45')
FORM_TYPES = {
    value.strip()
    for value in os.getenv('CPV45_FORM_TYPES', 'competition,change').split(',')
    if value.strip()
}
LATEST_VERSION_ONLY = env_bool('CPV45_LATEST_VERSION_ONLY', True)
INCLUDE_NOTICE_LEVEL_FALLBACK = env_bool('CPV45_NOTICE_FALLBACK', True)
REFRESH_EXPORTS = env_bool('CPV45_REFRESH_EXPORTS', False)

# Safe attachment policy.
DOWNLOAD_ATTACHMENTS = env_bool('CPV45_DOWNLOAD_ATTACHMENTS', False)
FOLLOW_STATIC_FILE_LINKS = env_bool('CPV45_FOLLOW_STATIC_LINKS', False)
REQUEST_DELAY_SECONDS = float(os.getenv('CPV45_REQUEST_DELAY_SECONDS', '1.0'))
MAX_ATTACHMENT_URLS = int(os.getenv('CPV45_MAX_ATTACHMENT_URLS', '50'))
MAX_FILES_PER_LANDING_PAGE = int(os.getenv('CPV45_MAX_FILES_PER_PAGE', '20'))
MAX_DOWNLOAD_BYTES = int(os.getenv('CPV45_MAX_DOWNLOAD_BYTES', str(100 * 1024 * 1024)))
MAX_ARCHIVE_BYTES = int(os.getenv('CPV45_MAX_ARCHIVE_BYTES', str(500 * 1024 * 1024)))
MAX_TEXT_CHARS_PER_FILE = int(os.getenv('CPV45_MAX_TEXT_PER_FILE', '100000'))
MAX_TEXT_CHARS_PER_LOT = int(os.getenv('CPV45_MAX_TEXT_PER_LOT', '300000'))

USER_AGENT = os.getenv(
    'CPV45_USER_AGENT',
    'CPV45TenderResearch/0.1 (public-procurement analysis)',
)

CACHE_DIRECTORY.mkdir(parents=True, exist_ok=True)
OUTPUT_DIRECTORY.mkdir(parents=True, exist_ok=True)
ATTACHMENT_DIRECTORY.mkdir(parents=True, exist_ok=True)
MANUAL_ATTACHMENTS_DIRECTORY.mkdir(parents=True, exist_ok=True)

print({
    'sourceMode': SOURCE_MODE,
    'period': f'{PERIOD_TYPE}:{PERIOD_VALUE}',
    'cpvPrefix': CPV_PREFIX,
    'formTypes': sorted(FORM_TYPES),
    'downloadAttachments': DOWNLOAD_ATTACHMENTS,
    'outputDirectory': str(OUTPUT_DIRECTORY),
})


# 2. Download and load both exports
#
# The CSV and eForms requests must use the same publication period. Files are cached so rerunning the notebook does not repeatedly download unchanged exports.


REQUIRED_CSV_TABLES = {
    'notice.csv',
    'lot.csv',
    'classification.csv',
    'purpose.csv',
}


def validate_zip(path: Path) -> None:
    if not path.exists():
        raise FileNotFoundError(path)
    if not zipfile.is_zipfile(path):
        raise ValueError(f'Not a valid ZIP file: {path}')
    with zipfile.ZipFile(path) as archive:
        bad_member = archive.testzip()
        if bad_member:
            raise ValueError(f'Corrupt ZIP member {bad_member!r} in {path}')


def download_to_path(url: str, destination: Path, max_bytes: int) -> Path:
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = destination.with_suffix(destination.suffix + '.part')
    request = Request(url, headers={'User-Agent': USER_AGENT})

    try:
        with urlopen(request, timeout=180, context=SSL_CONTEXT) as response, temporary.open('wb') as output:
            expected = response.headers.get('Content-Length')
            if expected and int(expected) > max_bytes:
                raise ValueError(f'Download exceeds configured size limit: {expected} bytes')

            downloaded = 0
            while True:
                chunk = response.read(1024 * 1024)
                if not chunk:
                    break
                downloaded += len(chunk)
                if downloaded > max_bytes:
                    raise ValueError(f'Download exceeded {max_bytes:,} bytes')
                output.write(chunk)

        temporary.replace(destination)
    except Exception:
        temporary.unlink(missing_ok=True)
        raise

    return destination


def export_url(format_name: str) -> str:
    if PERIOD_TYPE not in {'day', 'month'}:
        raise ValueError("PERIOD_TYPE must be 'day' or 'month'.")
    period_parameter = 'pubDay' if PERIOD_TYPE == 'day' else 'pubMonth'
    return API_URL + '?' + urlencode({
        period_parameter: PERIOD_VALUE,
        'format': format_name,
    })


def resolve_export_paths() -> tuple[Path, Path]:
    if SOURCE_MODE == 'local':
        validate_zip(LOCAL_CSV_ZIP)
        validate_zip(LOCAL_EFORMS_ZIP)
        return LOCAL_CSV_ZIP, LOCAL_EFORMS_ZIP

    if SOURCE_MODE != 'api':
        raise ValueError("SOURCE_MODE must be 'api' or 'local'.")

    safe_period = re.sub(r'[^0-9A-Za-z_-]+', '_', PERIOD_VALUE)
    csv_path = CACHE_DIRECTORY / f'{PERIOD_TYPE}_{safe_period}_csv.zip'
    eforms_path = CACHE_DIRECTORY / f'{PERIOD_TYPE}_{safe_period}_eforms.zip'

    for format_name, path in [('csv.zip', csv_path), ('eforms.zip', eforms_path)]:
        if REFRESH_EXPORTS or not path.exists():
            print(f'Downloading {format_name} ...')
            download_to_path(export_url(format_name), path, MAX_ARCHIVE_BYTES)
        validate_zip(path)

    return csv_path, eforms_path


def read_csv_stream(stream) -> pd.DataFrame:
    return pd.read_csv(
        stream,
        dtype=str,
        keep_default_na=False,
        encoding='utf-8-sig',
    ).fillna('')


def load_csv_tables(path: Path) -> dict[str, pd.DataFrame]:
    tables: dict[str, pd.DataFrame] = {}
    with zipfile.ZipFile(path) as archive:
        for member in archive.namelist():
            filename = Path(member).name
            if filename.lower().endswith('.csv'):
                with archive.open(member) as stream:
                    tables[filename] = read_csv_stream(stream)
    missing = REQUIRED_CSV_TABLES.difference(tables)
    if missing:
        raise ValueError(f'Required CSV tables are missing: {sorted(missing)}')
    return tables


csv_zip_path, eforms_zip_path = resolve_export_paths()
tables = load_csv_tables(csv_zip_path)

table_inventory = pd.DataFrame([
    {'table': name, 'rows': len(frame), 'columns': len(frame.columns)}
    for name, frame in sorted(tables.items())
])

print(f'CSV export: {csv_zip_path}')
print(f'eForms export: {eforms_zip_path}')
display(table_inventory)


# 3. Filter CPV 45 at lot level
#
# Both the main CPV and comma-separated additional CPV codes are checked. Lot-level classification takes priority. Notice-level CPV is used only when the notice has no usable lot-level classification.


NOTICE_KEY = ['noticeIdentifier', 'noticeVersion']
LOT_KEY = NOTICE_KEY + ['lotIdentifier']


def classification_codes(row: pd.Series) -> list[str]:
    values: list[str] = []
    main_code = str(row.get('mainClassificationCode', '')).strip()
    if main_code:
        values.append(main_code)
    additional = str(row.get('additionalClassificationCodes', '')).strip()
    if additional:
        values.extend(code.strip() for code in additional.split(',') if code.strip())
    return list(dict.fromkeys(values))


def combine_code_lists(series: Iterable[list[str]]) -> list[str]:
    return sorted({code for values in series for code in values})


classification = tables['classification.csv'].copy()
lots = tables['lot.csv'].copy()
notices = tables['notice.csv'].copy()

classification['allClassificationCodes'] = classification.apply(
    classification_codes,
    axis=1,
)
classification['matchedCpvCodes'] = classification['allClassificationCodes'].apply(
    lambda codes: [code for code in codes if code.startswith(CPV_PREFIX)]
)
classification['isTargetCpv'] = classification['matchedCpvCodes'].str.len().gt(0)
classification['hasUsableClassification'] = classification[
    'allClassificationCodes'
].str.len().gt(0)

lot_classifications = classification[classification['lotIdentifier'].ne('')].copy()
notice_classifications = classification[classification['lotIdentifier'].eq('')].copy()

target_lot_rows = lot_classifications[lot_classifications['isTargetCpv']]
if target_lot_rows.empty:
    lot_matches = pd.DataFrame(columns=LOT_KEY + ['matchedCpvCodes', 'matchedVia'])
else:
    lot_matches = (
        target_lot_rows.groupby(LOT_KEY, as_index=False)
        .agg(matchedCpvCodes=('matchedCpvCodes', combine_code_lists))
    )
    lot_matches['matchedVia'] = 'lot'

fallback_matches = pd.DataFrame(columns=LOT_KEY + ['matchedCpvCodes', 'matchedVia'])

if INCLUDE_NOTICE_LEVEL_FALLBACK:
    usable_lot_notice_keys = lot_classifications[
        lot_classifications['hasUsableClassification']
    ][NOTICE_KEY].drop_duplicates()

    notice_target_rows = notice_classifications[notice_classifications['isTargetCpv']]
    if not notice_target_rows.empty:
        notice_matches = (
            notice_target_rows.groupby(NOTICE_KEY, as_index=False)
            .agg(matchedCpvCodes=('matchedCpvCodes', combine_code_lists))
        )
        notice_matches = notice_matches.merge(
            usable_lot_notice_keys.assign(_hasLotClassification=True),
            on=NOTICE_KEY,
            how='left',
        )
        notice_matches = notice_matches[
            notice_matches['_hasLotClassification'].isna()
        ].drop(columns='_hasLotClassification')

        fallback_matches = lots.merge(
            notice_matches,
            on=NOTICE_KEY,
            how='inner',
        )[LOT_KEY + ['matchedCpvCodes']]
        fallback_matches['matchedVia'] = 'notice_fallback'

matched_lots = pd.concat([lot_matches, fallback_matches], ignore_index=True)
matched_lots['_priority'] = matched_lots['matchedVia'].map({
    'lot': 0,
    'notice_fallback': 1,
})
matched_lots = (
    matched_lots.sort_values('_priority')
    .drop_duplicates(LOT_KEY, keep='first')
    .drop(columns='_priority')
)

if FORM_TYPES and 'formType' in notices.columns:
    allowed_notice_keys = notices[notices['formType'].isin(FORM_TYPES)][
        NOTICE_KEY
    ].drop_duplicates()
    matched_lots = matched_lots.merge(allowed_notice_keys, on=NOTICE_KEY, how='inner')

if LATEST_VERSION_ONLY and not matched_lots.empty:
    version_order = notices.copy()
    version_order['_versionNumber'] = pd.to_numeric(
        version_order['noticeVersion'], errors='coerce'
    ).fillna(-1)
    latest_notice_keys = (
        version_order.sort_values(['noticeIdentifier', '_versionNumber'])
        .groupby('noticeIdentifier', as_index=False)
        .tail(1)[NOTICE_KEY]
    )
    matched_lots = matched_lots.merge(latest_notice_keys, on=NOTICE_KEY, how='inner')


def build_candidate_summary() -> pd.DataFrame:
    summary = matched_lots.merge(notices, on=NOTICE_KEY, how='left')
    purpose = tables['purpose.csv'].copy()
    fields = [
        'internalIdentifier',
        'mainNature',
        'additionalNature',
        'title',
        'estimatedValue',
        'estimatedValueCurrency',
        'description',
    ]

    lot_purpose = purpose[purpose['lotIdentifier'].ne('')][LOT_KEY + fields].copy()
    lot_purpose = lot_purpose.drop_duplicates(LOT_KEY)
    lot_purpose = lot_purpose.rename(columns={field: f'lot_{field}' for field in fields})

    notice_purpose = purpose[purpose['lotIdentifier'].eq('')][NOTICE_KEY + fields].copy()
    notice_purpose = notice_purpose.drop_duplicates(NOTICE_KEY)
    notice_purpose = notice_purpose.rename(
        columns={field: f'notice_{field}' for field in fields}
    )

    summary = summary.merge(lot_purpose, on=LOT_KEY, how='left')
    summary = summary.merge(notice_purpose, on=NOTICE_KEY, how='left')

    for field in fields:
        lot_column = f'lot_{field}'
        notice_column = f'notice_{field}'
        summary[field] = summary[lot_column].where(
            summary[lot_column].fillna('').ne(''),
            summary[notice_column],
        )

    summary = summary.drop(columns=[
        column
        for field in fields
        for column in (f'lot_{field}', f'notice_{field}')
    ])

    preferred = LOT_KEY + [
        'procedureIdentifier',
        'formType',
        'noticeType',
        'publicationDate',
        'matchedVia',
        'matchedCpvCodes',
        'internalIdentifier',
        'title',
        'description',
        'mainNature',
        'estimatedValue',
        'estimatedValueCurrency',
    ]
    return summary[[column for column in preferred if column in summary.columns]].copy()


candidate_summary = build_candidate_summary()
candidate_export = candidate_summary.copy()
candidate_export['matchedCpvCodes'] = candidate_export['matchedCpvCodes'].apply(
    lambda value: json.dumps(value, ensure_ascii=False)
)
candidate_path = OUTPUT_DIRECTORY / 'cpv45_candidates.csv'
candidate_export.to_csv(candidate_path, index=False, encoding='utf-8-sig')

print(f'Matching notices: {len(matched_lots[NOTICE_KEY].drop_duplicates()):,}')
print(f'Matching lots: {len(matched_lots):,}')
print(f'Saved: {candidate_path}')
display(candidate_export.head(10))


# 4. Parse the original eForms XML
#
# The parser extracts the notice and lot fields needed for early bid/no-bid decisions. Lists are retained as structured JSON rather than flattened into a single prompt.


NS = {
    'cac': 'urn:oasis:names:specification:ubl:schema:xsd:CommonAggregateComponents-2',
    'cbc': 'urn:oasis:names:specification:ubl:schema:xsd:CommonBasicComponents-2',
    'efac': 'http://data.europa.eu/p27/eforms-ubl-extension-aggregate-components/1',
    'efbc': 'http://data.europa.eu/p27/eforms-ubl-extension-basic-components/1',
}


def clean_text(value: str | None) -> str:
    return re.sub(r'\s+', ' ', value or '').strip()


def element_text(parent: ET.Element | None, path: str) -> str:
    if parent is None:
        return ''
    return clean_text(parent.findtext(path, default='', namespaces=NS))


def unique_strings(values: Iterable[str]) -> list[str]:
    return list(dict.fromkeys(value for value in values if value))


def normalize_version(value: Any) -> str:
    text = str(value or '').strip()
    if text.isdigit():
        return str(int(text))
    return text


def element_list_name(parent: ET.Element, path: str) -> str:
    element = parent.find(path, NS)
    return element.attrib.get('listName', '') if element is not None else ''


def parse_organizations(root: ET.Element) -> dict[str, str]:
    organizations: dict[str, str] = {}
    for organization in root.findall('.//efac:Organization', NS):
        identifier = element_text(
            organization,
            './efac:Company/cac:PartyIdentification/cbc:ID',
        )
        name = element_text(organization, './efac:Company/cac:PartyName/cbc:Name')
        if identifier:
            organizations[identifier] = name
    return organizations


def parse_project(project: ET.Element | None) -> dict[str, Any]:
    if project is None:
        return {}

    amount_element = project.find(
        './cac:RequestedTenderTotal/cbc:EstimatedOverallContractAmount',
        NS,
    )
    amount = clean_text(amount_element.text) if amount_element is not None else ''
    currency = amount_element.attrib.get('currencyID', '') if amount_element is not None else ''

    cpv_codes = unique_strings([
        element_text(project, './cac:MainCommodityClassification/cbc:ItemClassificationCode'),
        *[
            clean_text(element.text)
            for element in project.findall(
                './cac:AdditionalCommodityClassification/cbc:ItemClassificationCode',
                NS,
            )
        ],
    ])

    locations = []
    for location in project.findall('./cac:RealizedLocation/cac:Address', NS):
        locations.append({
            'street': element_text(location, './cbc:StreetName'),
            'additionalStreet': element_text(location, './cbc:AdditionalStreetName'),
            'postcode': element_text(location, './cbc:PostalZone'),
            'city': element_text(location, './cbc:CityName'),
            'nuts': element_text(location, './cbc:CountrySubentityCode'),
            'country': element_text(location, './cac:Country/cbc:IdentificationCode'),
        })

    return {
        'internalIdentifier': element_text(project, './cbc:ID'),
        'title': element_text(project, './cbc:Name'),
        'description': element_text(project, './cbc:Description'),
        'contractNature': element_text(project, './cbc:ProcurementTypeCode'),
        'estimatedValue': amount,
        'estimatedValueCurrency': currency,
        'cpvCodes': cpv_codes,
        'locations': locations,
        'durationStartDate': element_text(project, './cac:PlannedPeriod/cbc:StartDate'),
        'durationEndDate': element_text(project, './cac:PlannedPeriod/cbc:EndDate'),
        'durationMeasure': element_text(project, './cac:PlannedPeriod/cbc:DurationMeasure'),
    }


def parse_requirements(terms: ET.Element | None) -> dict[str, list[dict[str, str]]]:
    result: dict[str, list[dict[str, str]]] = {
        'selectionCriteria': [],
        'qualificationRequirements': [],
        'guarantees': [],
        'executionConditions': [],
    }
    if terms is None:
        return result

    for criterion in terms.findall('.//efac:SelectionCriteria', NS):
        result['selectionCriteria'].append({
            'code': element_text(criterion, './cbc:TendererRequirementTypeCode'),
            'description': element_text(criterion, './cbc:Description'),
        })

    for requirement in terms.findall(
        './/cac:TendererQualificationRequest/cac:SpecificTendererRequirement',
        NS,
    ):
        code_element = requirement.find('./cbc:TendererRequirementTypeCode', NS)
        result['qualificationRequirements'].append({
            'code': clean_text(code_element.text) if code_element is not None else '',
            'codeList': code_element.attrib.get('listName', '') if code_element is not None else '',
            'description': element_text(requirement, './cbc:Description'),
        })

    for guarantee in terms.findall('.//cac:RequiredFinancialGuarantee', NS):
        result['guarantees'].append({
            'required': element_text(guarantee, './cbc:GuaranteeTypeCode'),
            'description': element_text(guarantee, './cbc:Description'),
        })

    for condition in terms.findall('.//cac:ContractExecutionRequirement', NS):
        code_element = condition.find('./cbc:ExecutionRequirementCode', NS)
        result['executionConditions'].append({
            'code': clean_text(code_element.text) if code_element is not None else '',
            'codeList': code_element.attrib.get('listName', '') if code_element is not None else '',
            'description': element_text(condition, './cbc:Description'),
        })

    return result


def normalize_document_url(value: str) -> str:
    value = value.strip()
    if not value:
        return ''
    parsed = urlparse(value)
    if not parsed.scheme and '.' in value.split('/')[0]:
        return 'https://' + value
    return value


def parse_document_links(terms: ET.Element | None, scope: str) -> list[dict[str, str]]:
    if terms is None:
        return []
    links = []
    for reference in terms.findall('.//cac:CallForTendersDocumentReference', NS):
        raw_url = element_text(
            reference,
            './cac:Attachment/cac:ExternalReference/cbc:URI',
        )
        if raw_url:
            links.append({
                'scope': scope,
                'documentReferenceId': element_text(reference, './cbc:ID'),
                'documentType': element_text(reference, './cbc:DocumentType'),
                'description': element_text(reference, './cbc:DocumentDescription'),
                'rawUrl': raw_url,
                'url': normalize_document_url(raw_url),
            })
    return links


def parse_deadlines(process: ET.Element | None) -> dict[str, str]:
    if process is None:
        return {}
    return {
        'submissionDeadlineDate': element_text(
            process,
            './cac:TenderSubmissionDeadlinePeriod/cbc:EndDate',
        ),
        'submissionDeadlineTime': element_text(
            process,
            './cac:TenderSubmissionDeadlinePeriod/cbc:EndTime',
        ),
        'questionDeadlineDate': element_text(
            process,
            './cac:AdditionalInformationRequestPeriod/cbc:EndDate',
        ),
        'questionDeadlineTime': element_text(
            process,
            './cac:AdditionalInformationRequestPeriod/cbc:EndTime',
        ),
        'submissionMethod': element_text(process, './cbc:SubmissionMethodCode'),
        'submissionUrl': element_text(process, './cbc:AccessToolsURI'),
        'procedureType': element_text(process, './cbc:ProcedureCode'),
    }


def prefer_value(primary: Any, fallback: Any) -> Any:
    return primary if primary not in ('', None, [], {}) else fallback


def merge_lists(primary: list[Any], fallback: list[Any]) -> list[Any]:
    values = primary if primary else fallback
    seen: set[str] = set()
    result = []
    for value in values:
        marker = json.dumps(value, ensure_ascii=False, sort_keys=True)
        if marker not in seen:
            seen.add(marker)
            result.append(value)
    return result


def parse_eforms_xml(content: bytes, filename: str) -> list[dict[str, Any]]:
    root = ET.fromstring(content)
    notice_id = element_text(root, "./cbc:ID[@schemeName='notice-id']")
    if not notice_id:
        notice_id = element_text(root, './cbc:ID')
    notice_version = element_text(root, './cbc:VersionID')

    notice_type_element = root.find('./cbc:NoticeTypeCode', NS)
    notice_type_code = clean_text(notice_type_element.text) if notice_type_element is not None else ''
    notice_type_group = (
        notice_type_element.attrib.get('listName', '')
        if notice_type_element is not None
        else ''
    )

    organizations = parse_organizations(root)
    buyer_references = [
        clean_text(element.text)
        for element in root.findall(
            './cac:ContractingParty/cac:Party/cac:PartyIdentification/cbc:ID',
            NS,
        )
    ]
    buyer_names = unique_strings([
        organizations.get(reference, reference)
        for reference in buyer_references
    ])

    notice_project = parse_project(root.find('./cac:ProcurementProject', NS))
    notice_terms = root.find('./cac:TenderingTerms', NS)
    notice_process = root.find('./cac:TenderingProcess', NS)
    notice_requirements = parse_requirements(notice_terms)
    notice_documents = parse_document_links(notice_terms, 'notice')
    notice_deadlines = parse_deadlines(notice_process)

    common = {
        'noticeIdentifier': notice_id,
        'noticeVersion': notice_version,
        'noticeVersionKey': normalize_version(notice_version),
        'procedureIdentifier': element_text(root, './cbc:ContractFolderID'),
        'noticeRootType': root.tag.split('}')[-1],
        'noticeTypeCode': notice_type_code,
        'noticeTypeGroup': notice_type_group,
        'customizationId': element_text(root, './cbc:CustomizationID'),
        'profileId': element_text(root, './cbc:ProfileID'),
        'issueDate': element_text(root, './cbc:IssueDate'),
        'requestedPublicationDate': element_text(root, './cbc:RequestedPublicationDate'),
        'buyerNames': buyer_names,
        'xmlFile': filename,
    }

    records: list[dict[str, Any]] = []
    lots_in_xml = root.findall('./cac:ProcurementProjectLot', NS)

    if not lots_in_xml:
        records.append({
            **common,
            'lotIdentifier': '',
            **notice_project,
            **notice_deadlines,
            **notice_requirements,
            'documentLinks': notice_documents,
            'awardCriterionTypes': unique_strings([
                clean_text(element.text)
                for element in root.findall(
                    './/cbc:AwardingCriterionTypeCode',
                    NS,
                )
            ]),
        })
        return records

    for lot in lots_in_xml:
        lot_project = parse_project(lot.find('./cac:ProcurementProject', NS))
        lot_terms = lot.find('./cac:TenderingTerms', NS)
        lot_process = lot.find('./cac:TenderingProcess', NS)
        lot_requirements = parse_requirements(lot_terms)
        lot_documents = parse_document_links(lot_terms, 'lot')
        lot_deadlines = parse_deadlines(lot_process)

        project = {
            key: prefer_value(lot_project.get(key), notice_project.get(key))
            for key in set(notice_project) | set(lot_project)
        }
        deadlines = {
            key: prefer_value(lot_deadlines.get(key), notice_deadlines.get(key))
            for key in set(notice_deadlines) | set(lot_deadlines)
        }
        requirements = {
            key: merge_lists(lot_requirements.get(key, []), notice_requirements.get(key, []))
            for key in notice_requirements
        }

        records.append({
            **common,
            'lotIdentifier': element_text(lot, './cbc:ID'),
            **project,
            **deadlines,
            **requirements,
            'documentLinks': merge_lists(lot_documents, notice_documents),
            'awardCriterionTypes': unique_strings([
                clean_text(element.text)
                for element in lot.findall('.//cbc:AwardingCriterionTypeCode', NS)
            ]),
        })

    return records


eforms_records: list[dict[str, Any]] = []
eforms_parse_errors: list[dict[str, str]] = []

with zipfile.ZipFile(eforms_zip_path) as archive:
    xml_members = [
        member for member in archive.namelist()
        if member.lower().endswith('.xml')
    ]
    for member in xml_members:
        try:
            eforms_records.extend(parse_eforms_xml(archive.read(member), member))
        except Exception as error:
            eforms_parse_errors.append({
                'xmlFile': member,
                'error': f'{type(error).__name__}: {error}',
            })

candidate_key_set = {
    (
        str(row.noticeIdentifier),
        normalize_version(row.noticeVersion),
        str(row.lotIdentifier),
    )
    for row in matched_lots.itertuples(index=False)
}

selected_eforms_records = [
    record
    for record in eforms_records
    if (
        record['noticeIdentifier'],
        record['noticeVersionKey'],
        record['lotIdentifier'],
    ) in candidate_key_set
]


def csv_safe_record(record: dict[str, Any]) -> dict[str, Any]:
    return {
        key: (
            json.dumps(value, ensure_ascii=False)
            if isinstance(value, (list, dict))
            else value
        )
        for key, value in record.items()
    }


eforms_details_path = OUTPUT_DIRECTORY / 'eforms_lot_details.csv'
pd.DataFrame([
    csv_safe_record(record) for record in selected_eforms_records
]).to_csv(eforms_details_path, index=False, encoding='utf-8-sig')

parse_errors_path = OUTPUT_DIRECTORY / 'eforms_parse_errors.csv'
pd.DataFrame(eforms_parse_errors, columns=['xmlFile', 'error']).to_csv(
    parse_errors_path,
    index=False,
    encoding='utf-8-sig',
)

document_link_rows: list[dict[str, str]] = []
for record in selected_eforms_records:
    for link in record.get('documentLinks', []):
        document_link_rows.append({
            'noticeIdentifier': record['noticeIdentifier'],
            'noticeVersion': record['noticeVersion'],
            'lotIdentifier': record['lotIdentifier'],
            **link,
        })

document_links = pd.DataFrame(
    document_link_rows,
    columns=LOT_KEY + [
        'scope',
        'documentReferenceId',
        'documentType',
        'description',
        'rawUrl',
        'url',
    ],
).drop_duplicates()

document_links_path = OUTPUT_DIRECTORY / 'document_links.csv'
document_links.to_csv(document_links_path, index=False, encoding='utf-8-sig')

print(f'XML files parsed: {len(xml_members):,}')
print(f'eForms lot records matched: {len(selected_eforms_records):,}')
print(f'Document links: {len(document_links):,}')
print(f'XML parse errors: {len(eforms_parse_errors):,}')
print(f'Saved: {eforms_details_path}')
print(f'Saved: {document_links_path}')
display(document_links.head(10))


# 5. Safely handle tender-document URLs
#
# A document URL may be either a direct file or an HTML landing page on RIB, DTVP, eVergabe or another platform. The generic downloader saves direct files. It does not attempt to bypass authentication, JavaScript challenges, CAPTCHAs or platform access controls.
#
# When a row receives `landing_page_requires_adapter`, open that URL manually or add a platform-specific adapter. You can place manually downloaded files under:
#
# `data/manual_attachments/<noticeIdentifier>/<lotIdentifier>/filename.pdf`


KNOWN_FILE_EXTENSIONS = {
    '.pdf', '.doc', '.docx', '.xls', '.xlsx', '.xlsm',
    '.zip', '.7z', '.rar', '.txt', '.csv', '.xml', '.json',
    '.x83', '.d83', '.p83', '.gaeb',
}


class LinkCollector(HTMLParser):
    def __init__(self):
        super().__init__()
        self.links: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag.lower() != 'a':
            return
        attributes = dict(attrs)
        href = attributes.get('href')
        if href:
            self.links.append(href)


def safe_component(value: str, fallback: str = 'unknown') -> str:
    cleaned = re.sub(r'[^0-9A-Za-z._-]+', '_', str(value)).strip('._')
    return cleaned[:120] or fallback


def read_response_limited(response, max_bytes: int) -> bytes:
    data = response.read(max_bytes + 1)
    if len(data) > max_bytes:
        raise ValueError(f'Response exceeded {max_bytes:,} bytes')
    return data


def response_filename(url: str, headers) -> str:
    disposition = headers.get('Content-Disposition', '')
    utf_match = re.search(r"filename\*=UTF-8''([^;]+)", disposition, flags=re.I)
    plain_match = re.search(r'filename="?([^";]+)', disposition, flags=re.I)
    if utf_match:
        name = unquote(utf_match.group(1))
    elif plain_match:
        name = plain_match.group(1).strip()
    else:
        name = Path(unquote(urlparse(url).path)).name
    return safe_component(name, 'attachment.bin')


def looks_like_direct_file(url: str, content_type: str, disposition: str) -> bool:
    extension = Path(urlparse(url).path).suffix.lower()
    if 'attachment' in disposition.lower():
        return True
    if extension in KNOWN_FILE_EXTENSIONS:
        return True
    if content_type.startswith('text/html') or content_type.startswith('application/json'):
        return False
    return bool(content_type)


def save_downloaded_file(
    content: bytes,
    final_url: str,
    headers,
    target_directory: Path,
) -> Path:
    target_directory.mkdir(parents=True, exist_ok=True)
    filename = response_filename(final_url, headers)
    path = target_directory / filename
    if path.exists() and hashlib.sha256(path.read_bytes()).digest() == hashlib.sha256(content).digest():
        return path
    if path.exists():
        path = target_directory / f'{path.stem}_{hashlib.sha256(content).hexdigest()[:8]}{path.suffix}'
    path.write_bytes(content)
    return path


def fetch_document_url(url: str, target_directory: Path) -> list[dict[str, Any]]:
    request = Request(url, headers={'User-Agent': USER_AGENT})
    with urlopen(request, timeout=90, context=SSL_CONTEXT) as response:
        final_url = response.geturl()
        content_type = response.headers.get_content_type()
        disposition = response.headers.get('Content-Disposition', '')
        content = read_response_limited(response, MAX_DOWNLOAD_BYTES)
        headers = response.headers

    if looks_like_direct_file(final_url, content_type, disposition):
        path = save_downloaded_file(content, final_url, headers, target_directory)
        return [{
            'status': 'downloaded_direct_file',
            'sourceUrl': url,
            'finalUrl': final_url,
            'contentType': content_type,
            'localPath': str(path),
            'sha256': hashlib.sha256(content).hexdigest(),
            'sizeBytes': len(content),
        }]

    results: list[dict[str, Any]] = [{
        'status': 'landing_page_requires_adapter',
        'sourceUrl': url,
        'finalUrl': final_url,
        'contentType': content_type,
        'localPath': '',
        'sha256': '',
        'sizeBytes': len(content),
    }]

    if not FOLLOW_STATIC_FILE_LINKS:
        return results

    parser = LinkCollector()
    parser.feed(content.decode('utf-8', errors='replace'))
    final_host = urlparse(final_url).hostname
    candidate_urls = []
    for href in parser.links:
        candidate_url = urljoin(final_url, href)
        parsed = urlparse(candidate_url)
        if parsed.scheme not in {'http', 'https'}:
            continue
        if parsed.hostname != final_host:
            continue
        if Path(parsed.path).suffix.lower() in KNOWN_FILE_EXTENSIONS:
            candidate_urls.append(candidate_url)

    for candidate_url in unique_strings(candidate_urls)[:MAX_FILES_PER_LANDING_PAGE]:
        time.sleep(REQUEST_DELAY_SECONDS)
        try:
            request = Request(candidate_url, headers={'User-Agent': USER_AGENT})
            with urlopen(request, timeout=90, context=SSL_CONTEXT) as response:
                child_type = response.headers.get_content_type()
                child_disposition = response.headers.get('Content-Disposition', '')
                child_content = read_response_limited(response, MAX_DOWNLOAD_BYTES)
                child_final_url = response.geturl()
                child_headers = response.headers
            if not looks_like_direct_file(child_final_url, child_type, child_disposition):
                continue
            path = save_downloaded_file(
                child_content,
                child_final_url,
                child_headers,
                target_directory,
            )
            results.append({
                'status': 'downloaded_static_page_link',
                'sourceUrl': candidate_url,
                'finalUrl': child_final_url,
                'contentType': child_type,
                'localPath': str(path),
                'sha256': hashlib.sha256(child_content).hexdigest(),
                'sizeBytes': len(child_content),
            })
        except Exception as error:
            results.append({
                'status': 'static_link_download_failed',
                'sourceUrl': candidate_url,
                'finalUrl': '',
                'contentType': '',
                'localPath': '',
                'sha256': '',
                'sizeBytes': 0,
                'error': f'{type(error).__name__}: {error}',
            })

    return results


attachment_manifest_rows: list[dict[str, Any]] = []

if not DOWNLOAD_ATTACHMENTS:
    for row in document_links.to_dict('records'):
        attachment_manifest_rows.append({
            **{key: row.get(key, '') for key in LOT_KEY},
            'documentType': row.get('documentType', ''),
            'documentUrl': row.get('url', ''),
            'status': 'not_attempted_download_disabled',
            'sourceUrl': row.get('url', ''),
            'finalUrl': '',
            'contentType': '',
            'localPath': '',
            'sha256': '',
            'sizeBytes': 0,
            'error': '',
        })
else:
    for index, row in enumerate(document_links.head(MAX_ATTACHMENT_URLS).to_dict('records')):
        if index:
            time.sleep(REQUEST_DELAY_SECONDS)
        if row.get('documentType') == 'restricted-document':
            attachment_manifest_rows.append({
                **{key: row.get(key, '') for key in LOT_KEY},
                'documentType': row.get('documentType', ''),
                'documentUrl': row.get('url', ''),
                'status': 'restricted_document_manual_access_required',
                'sourceUrl': row.get('url', ''),
                'finalUrl': '',
                'contentType': '',
                'localPath': '',
                'sha256': '',
                'sizeBytes': 0,
                'error': '',
            })
            continue
        target = (
            ATTACHMENT_DIRECTORY
            / safe_component(row['noticeIdentifier'])
            / safe_component(row['noticeVersion'])
            / safe_component(row['lotIdentifier'], '_notice')
        )
        try:
            download_results = fetch_document_url(row['url'], target)
            for result in download_results:
                attachment_manifest_rows.append({
                    **{key: row.get(key, '') for key in LOT_KEY},
                    'documentType': row.get('documentType', ''),
                    'documentUrl': row.get('url', ''),
                    'error': '',
                    **result,
                })
        except Exception as error:
            attachment_manifest_rows.append({
                **{key: row.get(key, '') for key in LOT_KEY},
                'documentType': row.get('documentType', ''),
                'documentUrl': row.get('url', ''),
                'status': 'download_failed',
                'sourceUrl': row.get('url', ''),
                'finalUrl': '',
                'contentType': '',
                'localPath': '',
                'sha256': '',
                'sizeBytes': 0,
                'error': f'{type(error).__name__}: {error}',
            })

attachment_manifest = pd.DataFrame(attachment_manifest_rows)
attachment_manifest_path = OUTPUT_DIRECTORY / 'attachment_download_manifest.csv'
attachment_manifest.to_csv(attachment_manifest_path, index=False, encoding='utf-8-sig')

print(f'Attachment download enabled: {DOWNLOAD_ATTACHMENTS}')
print(f'Saved: {attachment_manifest_path}')
if not attachment_manifest.empty:
    display(attachment_manifest['status'].value_counts().rename_axis('status').reset_index(name='count'))


# 6. Extract text from downloaded and manually supplied files
#
# Optional parsers improve coverage:
#
# ```python
# %pip install pypdf python-docx openpyxl
# ```
#
# Scanned PDFs still require an OCR component. They are marked with `needsOcr = true` when little or no text is found.


try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

try:
    from docx import Document
except ImportError:
    Document = None

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None


TEXT_EXTENSIONS = {
    '.txt', '.csv', '.tsv', '.xml', '.json', '.html', '.htm', '.md',
    '.x83', '.d83', '.p83', '.gaeb',
}


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open('rb') as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b''):
            digest.update(chunk)
    return digest.hexdigest()


def safe_extract_zip(path: Path, destination: Path) -> list[Path]:
    destination.mkdir(parents=True, exist_ok=True)
    root = destination.resolve()
    extracted: list[Path] = []
    total_size = 0

    with zipfile.ZipFile(path) as archive:
        if len(archive.infolist()) > 500:
            raise ValueError('Archive has more than 500 members')
        for info in archive.infolist():
            if info.is_dir():
                continue
            total_size += info.file_size
            if total_size > MAX_ARCHIVE_BYTES:
                raise ValueError('Uncompressed archive exceeds configured limit')
            target = (destination / info.filename).resolve()
            if root not in target.parents:
                raise ValueError(f'Unsafe archive path: {info.filename}')
            target.parent.mkdir(parents=True, exist_ok=True)
            with archive.open(info) as source, target.open('wb') as output:
                shutil.copyfileobj(source, output)
            extracted.append(target)

    return extracted


def extract_file_text(path: Path) -> tuple[str, str, bool]:
    extension = path.suffix.lower()

    if extension in TEXT_EXTENSIONS:
        text = path.read_text(encoding='utf-8', errors='replace')
        return text[:MAX_TEXT_CHARS_PER_FILE], 'text_extracted', False

    if extension == '.pdf':
        if PdfReader is None:
            return '', 'parser_missing_pypdf', True
        reader = PdfReader(str(path))
        text = '\n\n'.join(page.extract_text() or '' for page in reader.pages)
        text = text[:MAX_TEXT_CHARS_PER_FILE]
        return text, 'text_extracted' if text.strip() else 'no_embedded_text', len(text.strip()) < 100

    if extension == '.docx':
        if Document is None:
            return '', 'parser_missing_python_docx', False
        document = Document(str(path))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append('\t'.join(cell.text for cell in row.cells))
        text = '\n'.join(parts)[:MAX_TEXT_CHARS_PER_FILE]
        return text, 'text_extracted', False

    if extension in {'.xlsx', '.xlsm'}:
        if load_workbook is None:
            return '', 'parser_missing_openpyxl', False
        workbook = load_workbook(path, read_only=True, data_only=True)
        parts = []
        for worksheet in workbook.worksheets:
            parts.append(f'## Sheet: {worksheet.title}')
            for row in worksheet.iter_rows(values_only=True):
                parts.append('\t'.join('' if value is None else str(value) for value in row))
                if sum(len(part) for part in parts) >= MAX_TEXT_CHARS_PER_FILE:
                    break
        text = '\n'.join(parts)[:MAX_TEXT_CHARS_PER_FILE]
        return text, 'text_extracted', False

    return '', 'unsupported_file_type', False


candidate_version_lookup: dict[tuple[str, str], str] = {}
for row in matched_lots.itertuples(index=False):
    candidate_version_lookup[(str(row.noticeIdentifier), str(row.lotIdentifier))] = str(
        row.noticeVersion
    )

attachment_sources: list[dict[str, Any]] = []
if not attachment_manifest.empty:
    for row in attachment_manifest.to_dict('records'):
        local_path = str(row.get('localPath', '') or '')
        if local_path and Path(local_path).is_file():
            attachment_sources.append({
                **{key: str(row.get(key, '') or '') for key in LOT_KEY},
                'sourceUrl': str(row.get('sourceUrl', '') or ''),
                'localPath': local_path,
                'sourceType': 'automatic_download',
            })

for path in MANUAL_ATTACHMENTS_DIRECTORY.rglob('*'):
    if not path.is_file() or any(part.startswith('.') for part in path.parts):
        continue
    relative = path.relative_to(MANUAL_ATTACHMENTS_DIRECTORY)
    if len(relative.parts) < 3:
        continue
    notice_id, lot_id = relative.parts[0], relative.parts[1]
    notice_version = candidate_version_lookup.get((notice_id, lot_id), '')
    attachment_sources.append({
        'noticeIdentifier': notice_id,
        'noticeVersion': notice_version,
        'lotIdentifier': lot_id,
        'sourceUrl': '',
        'localPath': str(path),
        'sourceType': 'manual_directory',
    })

expanded_sources: list[dict[str, Any]] = []
for source in attachment_sources:
    path = Path(source['localPath'])
    if path.suffix.lower() != '.zip':
        expanded_sources.append(source)
        continue
    destination = path.parent / f'{path.stem}_expanded'
    try:
        for extracted_path in safe_extract_zip(path, destination):
            expanded_sources.append({
                **source,
                'localPath': str(extracted_path),
                'containerPath': str(path),
            })
    except Exception as error:
        expanded_sources.append({
            **source,
            'extractionError': f'{type(error).__name__}: {error}',
        })

extracted_documents: list[dict[str, Any]] = []
seen_document_paths: set[str] = set()
for source in expanded_sources:
    path = Path(source['localPath'])
    resolved = str(path.resolve())
    if resolved in seen_document_paths or not path.is_file():
        continue
    seen_document_paths.add(resolved)
    try:
        text, status, needs_ocr = extract_file_text(path)
        error = source.get('extractionError', '')
    except Exception as exception:
        text = ''
        status = 'text_extraction_failed'
        needs_ocr = False
        error = f'{type(exception).__name__}: {exception}'

    extracted_documents.append({
        **{key: source.get(key, '') for key in LOT_KEY},
        'sourceType': source.get('sourceType', ''),
        'sourceUrl': source.get('sourceUrl', ''),
        'fileName': path.name,
        'localPath': str(path),
        'containerPath': source.get('containerPath', ''),
        'extension': path.suffix.lower(),
        'sizeBytes': path.stat().st_size,
        'sha256': sha256_file(path),
        'textStatus': status,
        'needsOcr': needs_ocr,
        'text': text,
        'error': error,
    })

extracted_documents_path = OUTPUT_DIRECTORY / 'extracted_documents.jsonl'
with extracted_documents_path.open('w', encoding='utf-8') as output:
    for document in extracted_documents:
        output.write(json.dumps(document, ensure_ascii=False) + '\n')

document_summary_path = OUTPUT_DIRECTORY / 'extracted_documents_summary.csv'
pd.DataFrame([
    {key: value for key, value in document.items() if key != 'text'}
    for document in extracted_documents
]).to_csv(document_summary_path, index=False, encoding='utf-8-sig')

print(f'Attachment files available: {len(attachment_sources):,}')
print(f'Documents processed: {len(extracted_documents):,}')
print(f'Saved: {extracted_documents_path}')


# 7. Build one agent input object per lot
#
# The JSONL output keeps structured notice evidence separate from extracted document text. Each row also reports missing evidence so the agent cannot silently treat absent data as a passed requirement.


def clean_json_value(value: Any) -> Any:
    if isinstance(value, dict):
        return {key: clean_json_value(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [clean_json_value(item) for item in value]
    if isinstance(value, (pd.Timestamp,)):
        return value.isoformat()
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    return value


def lot_key(notice_id: Any, notice_version: Any, lot_id: Any) -> tuple[str, str, str]:
    return (
        str(notice_id),
        normalize_version(notice_version),
        str(lot_id),
    )


eforms_index = {
    lot_key(
        record['noticeIdentifier'],
        record['noticeVersion'],
        record['lotIdentifier'],
    ): record
    for record in selected_eforms_records
}

links_by_key: dict[tuple[str, str, str], list[dict[str, Any]]] = defaultdict(list)
for row in document_links.to_dict('records'):
    links_by_key[lot_key(row['noticeIdentifier'], row['noticeVersion'], row['lotIdentifier'])].append(
        clean_json_value(row)
    )

documents_by_key: dict[tuple[str, str, str], list[dict[str, Any]]] = defaultdict(list)
for document in extracted_documents:
    documents_by_key[lot_key(
        document['noticeIdentifier'],
        document['noticeVersion'],
        document['lotIdentifier'],
    )].append(document)


def limited_document_payload(documents: list[dict[str, Any]]) -> list[dict[str, Any]]:
    remaining = MAX_TEXT_CHARS_PER_LOT
    payload = []
    for document in documents:
        text = str(document.get('text', ''))
        included_text = text[:remaining]
        remaining -= len(included_text)
        payload.append({
            key: clean_json_value(value)
            for key, value in document.items()
            if key != 'text'
        } | {
            'text': included_text,
            'textTruncatedForAgent': len(included_text) < len(text),
        })
        if remaining <= 0:
            break
    return payload


agent_records: list[dict[str, Any]] = []
for discovery in candidate_summary.to_dict('records'):
    key = lot_key(
        discovery['noticeIdentifier'],
        discovery['noticeVersion'],
        discovery['lotIdentifier'],
    )
    eforms = eforms_index.get(key)
    links = links_by_key.get(key, [])
    documents = documents_by_key.get(key, [])

    missing_evidence = []
    if not eforms:
        missing_evidence.append('eforms_lot_not_matched')
    else:
        if not eforms.get('submissionDeadlineDate'):
            missing_evidence.append('submission_deadline_missing')
        if not eforms.get('selectionCriteria'):
            missing_evidence.append('selection_criteria_missing')
        if not eforms.get('documentLinks'):
            missing_evidence.append('procurement_document_link_missing')
    if links and not documents:
        missing_evidence.append('attachments_not_downloaded_or_unavailable')
    if any(document.get('needsOcr') for document in documents):
        missing_evidence.append('ocr_required')

    agent_records.append({
        'candidateId': ':'.join(key),
        'noticeIdentifier': discovery['noticeIdentifier'],
        'noticeVersion': discovery['noticeVersion'],
        'lotIdentifier': discovery['lotIdentifier'],
        'discovery': clean_json_value(discovery),
        'eforms': clean_json_value(eforms),
        'documentLinks': clean_json_value(links),
        'documents': limited_document_payload(documents),
        'missingEvidence': missing_evidence,
        'agentPolicy': {
            'hardExclusionsBeforeRanking': True,
            'doNotTreatMissingEvidenceAsPass': True,
            'requireEvidenceCitation': True,
        },
    })

agent_input_path = OUTPUT_DIRECTORY / 'agent_input.jsonl'
with agent_input_path.open('w', encoding='utf-8') as output:
    for record in agent_records:
        output.write(json.dumps(record, ensure_ascii=False) + '\n')

agent_overview = pd.DataFrame([{
    'candidateId': record['candidateId'],
    'title': (record.get('eforms') or {}).get('title') or record['discovery'].get('title'),
    'submissionDeadlineDate': (record.get('eforms') or {}).get('submissionDeadlineDate'),
    'selectionCriteriaCount': len((record.get('eforms') or {}).get('selectionCriteria', [])),
    'guaranteeCount': len((record.get('eforms') or {}).get('guarantees', [])),
    'documentLinkCount': len(record['documentLinks']),
    'documentCount': len(record['documents']),
    'missingEvidence': ' | '.join(record['missingEvidence']),
} for record in agent_records])

agent_overview_path = OUTPUT_DIRECTORY / 'agent_input_overview.csv'
agent_overview.to_csv(agent_overview_path, index=False, encoding='utf-8-sig')

print(f'Agent records: {len(agent_records):,}')
print(f'Saved: {agent_input_path}')
print(f'Saved: {agent_overview_path}')
display(agent_overview.head(20))


# 8. Validation
#
# These checks protect the agent from duplicate lot records, unsupported CPV matches and malformed document links. A successful run does not mean every tender document was downloaded; inspect `missingEvidence` and `attachment_download_manifest.csv`.


duplicate_candidates = candidate_summary.duplicated(LOT_KEY).sum()
if duplicate_candidates:
    raise AssertionError(f'{duplicate_candidates} duplicate candidate lot keys found')

invalid_cpv_matches = matched_lots[
    ~matched_lots['matchedCpvCodes'].apply(
        lambda codes: any(str(code).startswith(CPV_PREFIX) for code in codes)
    )
]
if not invalid_cpv_matches.empty:
    raise AssertionError('A retained lot has no matching CPV prefix evidence')

invalid_document_urls = document_links[
    ~document_links['url'].apply(
        lambda value: urlparse(str(value)).scheme in {'http', 'https'}
    )
] if not document_links.empty else document_links
if not invalid_document_urls.empty:
    raise AssertionError('A document URL is not HTTP or HTTPS')

if len(agent_records) != len(candidate_summary):
    raise AssertionError('Agent JSONL does not contain exactly one record per candidate lot')

matched_eforms_keys = {
    lot_key(record['noticeIdentifier'], record['noticeVersion'], record['lotIdentifier'])
    for record in selected_eforms_records
}
unmatched_key_count = len(candidate_key_set.difference(matched_eforms_keys))

validation_summary = pd.DataFrame([
    {'check': 'candidate lot keys unique', 'result': 'PASS', 'count': len(candidate_summary)},
    {'check': f'CPV starts with {CPV_PREFIX}', 'result': 'PASS', 'count': len(matched_lots)},
    {'check': 'valid document URL schemes', 'result': 'PASS', 'count': len(document_links)},
    {'check': 'one agent record per lot', 'result': 'PASS', 'count': len(agent_records)},
    {
        'check': 'candidate lots without exact eForms match',
        'result': 'REVIEW' if unmatched_key_count else 'PASS',
        'count': unmatched_key_count,
    },
    {
        'check': 'XML parse errors',
        'result': 'REVIEW' if eforms_parse_errors else 'PASS',
        'count': len(eforms_parse_errors),
    },
])

display(validation_summary)
print('Pipeline completed successfully.')


# Output files
#
# - `cpv45_candidates.csv`: fast discovery shortlist from CSV.
# - `eforms_lot_details.csv`: structured notice and lot evidence.
# - `document_links.csv`: procurement-document landing pages or direct links.
# - `attachment_download_manifest.csv`: download status and provenance.
# - `extracted_documents.jsonl`: extracted text and file metadata.
# - `agent_input.jsonl`: one complete input object per selected lot.
# - `agent_input_overview.csv`: compact review table for humans.
#
# For production, retain every notice version and every downloaded document checksum. Correction notices can change deadlines or replace tender documents.
