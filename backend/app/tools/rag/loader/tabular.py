import csv
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any

import requests
from openpyxl import load_workbook

from app.tools.rag.loader.base import BaseLoader, LoaderResult, SourceContent

# openpyxl reads the modern XML-based Excel formats only — not the legacy
# binary .xls format, which would need a different library (xlrd) entirely.
_EXCEL_EXTENSIONS = {".xlsx", ".xlsm"}


class TabularLoader(BaseLoader):
    """Loads a CSV or Excel (.xlsx/.xlsm) document from a local file path or
    an http(s) URL, dispatching on file extension. Both formats render to
    the same "Row N: col: value | col: value" text, using the header row
    for column names, ready for chunking.

    For Excel files, pass `metadata={"sheet_name": "..."}` via kwargs to
    read a specific worksheet — defaults to the active sheet.
    """

    def load(self, source_content: SourceContent, **kwargs: Any) -> LoaderResult:
        source_ref = source_content.source_ref
        is_excel = self._is_excel(source_ref)

        if is_excel:
            csv_text = self._xlsx_to_csv(source_content, kwargs)
        elif source_content.is_url():
            csv_text = self._load_from_url(source_content.source, kwargs)
        elif source_content.path_exists():
            csv_text = self._load_from_file(source_content.source)
        else:
            csv_text = source_content.source

        return self._parse_csv(csv_text, source_ref, source_format="xlsx" if is_excel else "csv")

    def _is_excel(self, source_ref: str) -> bool:
        return Path(source_ref.split("?")[0]).suffix.lower() in _EXCEL_EXTENSIONS

    def _xlsx_to_csv(self, source_content: SourceContent, kwargs: dict[str, Any]) -> str:
        """Reads one worksheet and renders it as CSV text, so _parse_csv (and
        everything downstream, like CsvChunker) doesn't need to know or care
        whether the original file was a .csv or an Excel workbook."""
        if source_content.is_url():
            try:
                response = requests.get(source_content.source, timeout=30)
                response.raise_for_status()
            except Exception as e:
                raise ValueError(f"Error fetching Excel file from URL {source_content.source}: {e}") from e
            workbook_source: Any = BytesIO(response.content)
        elif source_content.path_exists():
            workbook_source = source_content.source
        else:
            raise FileNotFoundError(f"Excel file not found: {source_content.source}")

        try:
            workbook = load_workbook(workbook_source, read_only=True, data_only=True)
        except Exception as e:
            raise ValueError(f"Error reading Excel file {source_content.source}: {e}") from e

        sheet_name = kwargs.get("metadata", {}).get("sheet_name")
        sheet = workbook[sheet_name] if sheet_name else workbook.active

        output = StringIO()
        writer = csv.writer(output)
        for row in sheet.iter_rows(values_only=True):
            writer.writerow(["" if cell is None else cell for cell in row])
        return output.getvalue()

    def _load_from_url(self, url: str, kwargs: dict[str, Any]) -> str:
        headers = kwargs.get(
            "headers",
            {
                "Accept": "text/csv, application/csv, text/plain",
                "User-Agent": "Mozilla/5.0 (compatible; app TabularLoader)",
            },
        )

        try:
            response = requests.get(url, headers=headers, timeout=30)
            response.raise_for_status()
            return response.text
        except Exception as e:
            raise ValueError(f"Error fetching CSV from URL {url}: {e}") from e

    def _load_from_file(self, path: str) -> str:
        with open(path, encoding="utf-8") as file:
            return file.read()

    def _parse_csv(self, content: str, source_ref: str, source_format: str) -> LoaderResult:
        try:
            csv_reader = csv.DictReader(StringIO(content))

            text_parts = []
            headers = csv_reader.fieldnames

            if headers:
                text_parts.append("Headers: " + " | ".join(headers))
                text_parts.append("-" * 50)

                for row_num, row in enumerate(csv_reader, 1):
                    row_text = " | ".join([f"{k}: {v}" for k, v in row.items() if v])
                    text_parts.append(f"Row {row_num}: {row_text}")

            text = "\n".join(text_parts)

            metadata: dict[str, Any] = {
                "format": source_format,
                "columns": headers,
                "rows": len(text_parts) - 2 if headers else 0,
            }

        except Exception as e:
            text = content
            metadata = {"format": source_format, "parse_error": str(e)}

        return LoaderResult(
            content=text,
            source=source_ref,
            metadata=metadata,
            doc_id=self.generate_doc_id(source_ref=source_ref, content=text),
        )
